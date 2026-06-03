from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 创建文档
doc = Document()

# 设置中文字体
doc.styles['Normal'].font.name = '微软雅黑'
doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# 标题
title = doc.add_heading('RegPilot 使用手册', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 副标题
subtitle = doc.add_paragraph('产品合规法规分析系统')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].italic = True

doc.add_paragraph()

# 目录
doc.add_heading('目录', level=1)
sections = [
    '1. 系统简介',
    '2. 环境要求',
    '3. 安装步骤',
    '4. 启动程序',
    '5. 功能说明',
    '6. 常见问题',
]
for s in sections:
    doc.add_paragraph(s, style='List Bullet')

doc.add_page_break()

# 1. 系统简介
doc.add_heading('1. 系统简介', level=1)
doc.add_paragraph(
    'RegPilot 是一款面向自行车/电助力车产品出口美国市场的 AI 合规分析系统。'
    '系统基于内置的 50 条法规规则（覆盖 CPSC 1512、CPSIA、UL 2849、UL 2271 等法规），'
    '自动解析用户上传的 BOM 物料清单，进行产品名称标准化、产品分类、法规匹配、'
    '差距检测、风险评估，最终生成完整的合规分析报告。'
)

doc.add_paragraph('核心功能：')
features = [
    'BOM 物料清单解析与标准化',
    '自动法规匹配与合规检查',
    '4 维度合规评分（法规符合率/必需项缺失/风险等级/认证完整度）',
    '风险评估与整改建议',
    'AI 法规顾问（本地知识库驱动，无需外部 API）',
]
for f in features:
    doc.add_paragraph(f, style='List Bullet')

# 2. 环境要求
doc.add_heading('2. 环境要求', level=1)
doc.add_paragraph('2.1 操作系统')
doc.add_paragraph('macOS / Linux / Windows', style='List Bullet')

doc.add_paragraph('2.2 Python 环境')
doc.add_paragraph('Python 3.10 或更高版本', style='List Bullet')
doc.add_paragraph('Node.js 18 或更高版本', style='List Bullet')

doc.add_paragraph('2.3 依赖包')
doc.add_paragraph('后端依赖：')
deps = [
    'fastapi',
    'uvicorn',
    'python-multipart',
    'openpyxl',
    'sqlalchemy[asyncio]',
    'asyncpg',
    'pydantic-settings',
    'openai（可选，用于增强 AI 顾问）',
]
for d in deps:
    doc.add_paragraph(f'  - {d}', style='List Bullet')

doc.add_paragraph('前端依赖：')
doc.add_paragraph('  - Next.js 15', style='List Bullet')
doc.add_paragraph('  - React 18', style='List Bullet')
doc.add_paragraph('  - TypeScript', style='List Bullet')
doc.add_paragraph('  - TailwindCSS', style='List Bullet')

# 3. 安装步骤
doc.add_heading('3. 安装步骤', level=1)

doc.add_paragraph('3.1 克隆项目')
doc.add_paragraph('项目已位于：/Users/ttt/Desktop/RuleMinder', style='List Bullet')

doc.add_paragraph('3.2 安装后端依赖')
code1 = 'cd /Users/ttt/Desktop/RuleMinder/backend\npip3 install fastapi uvicorn python-multipart openpyxl sqlalchemy[asyncio] asyncpg pydantic-settings openai'
doc.add_paragraph(code1, style='Quote')

doc.add_paragraph('3.3 安装前端依赖')
code2 = 'cd /Users/ttt/Desktop/RuleMinder/frontend\nnpm install'
doc.add_paragraph(code2, style='Quote')

doc.add_paragraph('3.4 配置环境变量（可选）')
doc.add_paragraph('如需使用 OpenAI 增强 AI 顾问功能，创建 backend/.env 文件：', style='List Bullet')
env_config = 'OPENAI_API_KEY=sk-your-api-key\nOPENAI_MODEL=gpt-4o-mini'
doc.add_paragraph(env_config, style='Quote')
doc.add_paragraph('注：不配置 OPENAI_API_KEY 也可正常使用，AI 顾问将基于本地知识库回答。', style='List Bullet')

doc.add_page_break()

# 4. 启动程序
doc.add_heading('4. 启动程序', level=1)

doc.add_paragraph('4.1 启动后端服务')
code3 = 'cd /Users/ttt/Desktop/RuleMinder/backend\npython3 -m uvicorn app.main:app --reload --port 8000'
doc.add_paragraph(code3, style='Quote')
doc.add_paragraph('后端服务将在 http://localhost:8000 运行', style='List Bullet')
doc.add_paragraph('API 文档地址：http://localhost:8000/docs', style='List Bullet')

doc.add_paragraph('4.2 启动前端服务')
code4 = 'cd /Users/ttt/Desktop/RuleMinder/frontend\nnpm run dev'
doc.add_paragraph(code4, style='Quote')
doc.add_paragraph('前端服务将在 http://localhost:3000 运行', style='List Bullet')

doc.add_paragraph('4.3 访问系统')
doc.add_paragraph('浏览器打开：http://localhost:3000/check', style='List Bullet')

# 5. 功能说明
doc.add_heading('5. 功能说明', level=1)

doc.add_paragraph('5.1 首页（/）')
doc.add_paragraph('查看最近 6 条分析记录（评分 + 风险等级 + 日期）', style='List Bullet')
doc.add_paragraph('快速上传 BOM 文件入口', style='List Bullet')

doc.add_paragraph('5.2 产品检查（/check）')
doc.add_paragraph('核心合规检查页面，操作流程：', style='List Bullet')
doc.add_paragraph('1. 选择目标市场（美国/欧盟/日本等）', style='List Bullet')
doc.add_paragraph('2. 选择产品类型（自行车/儿童自行车/电助力车）', style='List Bullet')
doc.add_paragraph('3. 上传产品资料（BOM 必填，规格书/说明文档/图片可选）', style='List Bullet')
doc.add_paragraph('4. 点击"开始合规分析"，系统自动执行 7 步分析流程', style='List Bullet')

doc.add_paragraph('5.3 分析报告（/analysis/[id]）')
doc.add_paragraph('查看完整合规分析报告，包含：', style='List Bullet')
doc.add_paragraph('综合评分环形图（0-100 分）', style='List Bullet')
doc.add_paragraph('4 维度分解：法规符合率、必需项缺失、风险等级、认证完整度', style='List Bullet')
doc.add_paragraph('适用法规标签（CPSC 1512/CPSIA/UL 2849/UL 2271）', style='List Bullet')
doc.add_paragraph('法规检查结果表格（可展开查看详情）', style='List Bullet')
doc.add_paragraph('整改建议编号列表（按风险等级排序）', style='List Bullet')
doc.add_paragraph('BOM 物料清单表格', style='List Bullet')

doc.add_paragraph('5.4 法规库（/regulations）')
doc.add_paragraph('浏览 50 条法规规则', style='List Bullet')
doc.add_paragraph('按分类筛选（制动/反光/车铃/链罩/标签/电池/电机/化学/测试等）', style='List Bullet')
doc.add_paragraph('按法规来源筛选（CPSC 1512/CPSIA/UL 2849/UL 2271）', style='List Bullet')
doc.add_paragraph('关键词搜索', style='List Bullet')

doc.add_paragraph('5.5 AI 法规顾问（/advisor）')
doc.add_paragraph('对话式法规咨询界面', style='List Bullet')
doc.add_paragraph('基于本地 50 条法规知识库智能检索', style='List Bullet')
doc.add_paragraph('推荐问题快捷提问', style='List Bullet')
doc.add_paragraph('配置 OPENAI_API_KEY 后可升级为 GPT 增强回答', style='List Bullet')

doc.add_page_break()

# 6. 常见问题
doc.add_heading('6. 常见问题', level=1)

doc.add_paragraph('Q1: 后端启动失败，提示"ModuleNotFoundError"？')
doc.add_paragraph('A: 确保已安装所有后端依赖：', style='List Bullet')
doc.add_paragraph('   pip3 install fastapi uvicorn python-multipart openpyxl', style='List Bullet')

doc.add_paragraph('Q2: 前端启动失败，提示"module not found"？')
doc.add_paragraph('A: 确保已安装前端依赖：', style='List Bullet')
doc.add_paragraph('   cd frontend && npm install', style='List Bullet')

doc.add_paragraph('Q3: AI 顾问无法使用？')
doc.add_paragraph('A: 检查后端服务是否正常运行。AI 顾问基于本地知识库，无需 OpenAI API Key 即可使用。', style='List Bullet')

doc.add_paragraph('Q4: 上传 BOM 文件后分析失败？')
doc.add_paragraph('A: 确保 BOM 文件格式正确（.xlsx 或.csv），包含零部件名称、数量等基本信息。', style='List Bullet')

doc.add_paragraph('Q5: 如何配置数据库？')
doc.add_paragraph('A: MVP 版本使用内存存储，重启后数据清空。生产环境可配置 PostgreSQL：', style='List Bullet')
doc.add_paragraph('   1. 安装 PostgreSQL', style='List Bullet')
doc.add_paragraph('   2. 创建数据库：CREATE DATABASE regpilot;', style='List Bullet')
doc.add_paragraph('   3. 配置 DATABASE_URL 环境变量', style='List Bullet')

doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph('技术支持与反馈', style='List Bullet')
doc.add_paragraph('如有问题或建议，请联系开发团队。', style='List Bullet')

# 保存文档
doc.save('/Users/ttt/Desktop/RuleMinder/使用手册.docx')
print('使用手册.docx 已生成成功！')
