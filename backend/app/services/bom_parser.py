from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from app.schemas.bom import BOMComponent, ComponentSpec

DATA_DIR = Path(__file__).resolve().parent.parent.parent / "data"
MAPPINGS_DIR = DATA_DIR / "mappings"

_name_map: dict | None = None


def _load_name_map() -> dict:
    global _name_map
    if _name_map is None:
        with open(MAPPINGS_DIR / "component_name_map.json") as f:
            data = json.load(f)
        _name_map = {}
        for entry in data["mappings"]:
            key = entry["raw_name"].lower().strip()
            _name_map[key] = entry
    return _name_map


def parse_excel(file_path: str | Path) -> list[dict]:
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = rows[0]
    col_map = _detect_columns(header_row)

    items = []
    for row in rows[1:]:
        if not any(row):
            continue
        item = {}
        for field, col_idx in col_map.items():
            if col_idx is not None and col_idx < len(row):
                item[field] = row[col_idx]
            else:
                item[field] = None
        if item.get("raw_name"):
            items.append(item)

    wb.close()
    return items


def _detect_columns(header_row: tuple) -> dict:
    col_map = {
        "raw_name": None,
        "raw_spec": None,
        "quantity": None,
        "material": None,
        "part_number": None,
    }

    name_keywords = ["part", "component", "item", "description", "name", "零件", "部件", "名称", "描述"]
    spec_keywords = ["spec", "specification", "规格", "参数"]
    qty_keywords = ["qty", "quantity", "数量"]
    mat_keywords = ["material", "材料", "材质"]
    pn_keywords = ["part no", "part number", "p/n", "编号", "料号"]

    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        val = str(cell).lower().strip()
        if col_map["raw_name"] is None and any(k in val for k in name_keywords):
            col_map["raw_name"] = idx
        elif col_map["raw_spec"] is None and any(k in val for k in spec_keywords):
            col_map["raw_spec"] = idx
        elif col_map["quantity"] is None and any(k in val for k in qty_keywords):
            col_map["quantity"] = idx
        elif col_map["material"] is None and any(k in val for k in mat_keywords):
            col_map["material"] = idx
        elif col_map["part_number"] is None and any(k in val for k in pn_keywords):
            col_map["part_number"] = idx

    if col_map["raw_name"] is None and len(header_row) > 0:
        col_map["raw_name"] = 0

    return col_map


def normalize_component(raw_name: str, raw_spec: str | None = None) -> BOMComponent:
    name_map = _load_name_map()
    key = raw_name.lower().strip()

    match = name_map.get(key)

    if match:
        return BOMComponent(
            standard_name=match["standard_name"],
            display_name=match["standard_name"].replace("_", " ").title(),
            category=match["category"],
            sub_category=match["sub_category"],
            original_names=[raw_name],
            norm_confidence=match["confidence"],
        )

    best_match = None
    best_score = 0.0
    for map_key, entry in name_map.items():
        score = _similarity(key, map_key)
        if score > best_score:
            best_score = score
            best_match = entry

    if best_match and best_score >= 0.60:
        return BOMComponent(
            standard_name=best_match["standard_name"],
            display_name=best_match["standard_name"].replace("_", " ").title(),
            category=best_match["category"],
            sub_category=best_match["sub_category"],
            original_names=[raw_name],
            norm_confidence=best_score * best_match["confidence"],
        )

    return BOMComponent(
        standard_name="unknown",
        display_name=raw_name,
        category="Unknown",
        sub_category=None,
        original_names=[raw_name],
        norm_confidence=0.0,
    )


def _similarity(a: str, b: str) -> float:
    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.80
    set_a = set(a.split())
    set_b = set(b.split())
    if not set_a or not set_b:
        return 0.0
    intersection = set_a & set_b
    union = set_a | set_b
    return len(intersection) / len(union)


def parse_and_normalize(file_path: str | Path) -> list[BOMComponent]:
    raw_items = parse_excel(file_path)
    components = []
    for item in raw_items:
        comp = normalize_component(
            raw_name=str(item.get("raw_name", "")),
            raw_spec=str(item.get("raw_spec", "")) if item.get("raw_spec") else None,
        )
        comp.quantity = int(item.get("quantity", 1)) if item.get("quantity") else 1
        comp.material = str(item.get("material", "")) if item.get("material") else None
        components.append(comp)
    return components
