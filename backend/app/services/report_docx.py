"""Generate Word (.docx) compliance analysis reports."""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

from app.schemas.analysis import AnalysisResponse

from app.config import settings
REPORT_DIR = Path(settings.REPORT_DIR)
REPORT_DIR.mkdir(parents=True, exist_ok=True)

COLOR_PRIMARY = RGBColor(0x4F, 0x46, 0xE5)
COLOR_MET = RGBColor(0x16, 0x65, 0x34)
COLOR_NOT_MET = RGBColor(0xDC, 0x26, 0x26)
COLOR_PARTIAL = RGBColor(0xD9, 0x77, 0x06)
COLOR_NOT_ASSESSED = RGBColor(0x6B, 0x72, 0x80)
COLOR_DARK = RGBColor(0x1E, 0x29, 0x3B)
COLOR_BODY = RGBColor(0x33, 0x41, 0x55)
COLOR_LIGHT = RGBColor(0x64, 0x74, 0x8B)


def _set_cell_shading(cell, color_hex: str):
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear',
    })
    shading.append(shd)


def _add_styled_paragraph(doc, text, size=11, bold=False, color=COLOR_BODY, space_after=6, alignment=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if alignment:
        p.alignment = alignment
    return p


def _status_color(status: str) -> RGBColor:
    return {
        "MET": COLOR_MET,
        "NOT_MET": COLOR_NOT_MET,
        "PARTIALLY_MET": COLOR_PARTIAL,
        "NOT_ASSESSED": COLOR_NOT_ASSESSED,
        "EXEMPT": COLOR_LIGHT,
    }.get(status, COLOR_BODY)


def _status_cn(status: str) -> str:
    return {
        "MET": "合规",
        "NOT_MET": "不合规",
        "PARTIALLY_MET": "部分合规",
        "NOT_ASSESSED": "待验证",
        "EXEMPT": "豁免",
    }.get(status, status)


def _risk_cn(level: str) -> str:
    return {
        "Critical": "严重",
        "Major": "重要",
        "Minor": "轻微",
        "Info": "提示",
    }.get(level, level)


def _product_type_cn(ptype: str) -> str:
    return {
        "Bicycle": "自行车",
        "Kids Bicycle": "儿童自行车",
        "E-bike": "电助力车",
        "Kids E-bike": "儿童电助力车",
    }.get(ptype, ptype)


def generate_report(analysis: AnalysisResponse) -> str:
    """Generate a concise Word report and return the file path."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ── Title ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("合规审查报告")
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    p.paragraph_format.space_after = Pt(16)

    # ── Basic Info ──
    product_name = analysis.product_name or "未命名产品"
    product_type = _product_type_cn(analysis.product_type)
    created = analysis.created_at[:19].replace("T", " ") if analysis.created_at else datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    compliance_score = analysis.compliance_score or 0
    risk_score = analysis.risk_score or 0
    risk_level = _risk_cn(analysis.risk_level or "Medium")

    summary = analysis.summary
    status_order = {"NOT_MET": 0, "PARTIALLY_MET": 1, "NOT_ASSESSED": 2, "MET": 3, "EXEMPT": 4}
    sorted_findings = sorted(analysis.findings, key=lambda f: (status_order.get(f.status, 5), -f.risk_score))
    not_met = [f for f in sorted_findings if f.status == "NOT_MET"]
    partially_met = [f for f in sorted_findings if f.status == "PARTIALLY_MET"]
    not_assessed = [f for f in sorted_findings if f.status == "NOT_ASSESSED"]

    info_table = doc.add_table(rows=8, cols=2, style='Table Grid')
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("产品名称", product_name),
        ("产品类型", product_type),
        ("审查日期", created),
        ("报告编号", f"RP-{analysis.id[:8].upper()}"),
        ("合规分数", f"{compliance_score:.1f}"),
        ("风险分数", f"{risk_score:.1f}"),
        ("风险等级", risk_level),
        ("合规/部分合规/不合规/待验证",
         f"{summary.met_count}/{summary.partially_met_count}/{summary.not_met_count}/{summary.not_assessed_count}"),
    ]
    for i, (label, value) in enumerate(info_data):
        cell_l = info_table.cell(i, 0)
        cell_r = info_table.cell(i, 1)
        cell_l.text = label
        cell_r.text = str(value)
        for cell in [cell_l, cell_r]:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
        cell_l.paragraphs[0].runs[0].bold = True
        _set_cell_shading(cell_l, "F1F5F9")

    doc.add_paragraph()

    # ── Issues (only NOT_MET and PARTIALLY_MET) ─
    issues = not_met + partially_met
    if issues:
        _add_styled_paragraph(doc, "不合规项", size=13, bold=True, color=COLOR_PRIMARY, space_after=10)

        header_row = ["序号", "规则ID", "要求摘要", "法规来源", "状态", "风险"]
        issues_table = doc.add_table(rows=1 + len(issues), cols=len(header_row), style='Table Grid')
        issues_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for j, header in enumerate(header_row):
            cell = issues_table.cell(0, j)
            cell.text = header
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _set_cell_shading(cell, "4F46E5")

        for i, f in enumerate(issues):
            row_data = [
                str(i + 1),
                f.rule_id,
                f.requirement_summary[:50] + ("..." if len(f.requirement_summary) > 50 else ""),
                f"{f.regulation.regulation} {f.regulation.section}",
                _status_cn(f.status),
                _risk_cn(f.risk_level),
            ]
            for j, text in enumerate(row_data):
                cell = issues_table.cell(i + 1, j)
                cell.text = text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
                        if j == 4:
                            run.font.color.rgb = _status_color(f.status)
                            run.bold = True

            if f.status == "NOT_MET":
                for j in range(len(header_row)):
                    _set_cell_shading(issues_table.cell(i + 1, j), "FEF2F2")
            elif f.status == "PARTIALLY_MET":
                for j in range(len(header_row)):
                    _set_cell_shading(issues_table.cell(i + 1, j), "FFFBEB")

        doc.add_paragraph()

    # ── Pending Verifications ──
    verifications = [f for f in sorted_findings if f.pending_verifications and f.status != "NOT_MET"]
    if verifications:
        _add_styled_paragraph(doc, "待验证项", size=13, bold=True, color=COLOR_PARTIAL, space_after=10)
        for i, f in enumerate(verifications, 1):
            verifs = "；".join(f.pending_verifications)
            _add_styled_paragraph(
                doc,
                f"{i}. [{f.rule_id}] {f.requirement_summary}",
                size=10, bold=True, color=COLOR_DARK, space_after=2,
            )
            _add_styled_paragraph(doc, f"   需验证：{verifs}", size=9, color=COLOR_BODY, space_after=6)

    doc.add_paragraph()

    # ── Disclaimer ──
    _add_styled_paragraph(
        doc,
        "免责声明：本报告仅供参考，实际合规状态应以第三方实验室测试报告和法规原文为准。",
        size=8, color=COLOR_LIGHT, alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _add_styled_paragraph(
        doc,
        f"Generated by RegPilot | {datetime.now().strftime('%Y-%m-%d')}",
        size=7, color=COLOR_LIGHT, alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    # ── Save ──
    safe_name = (analysis.product_name or "未命名产品").replace(" ", "_").replace("/", "_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{safe_name}_合规审查报告_{timestamp}.docx"
    filepath = REPORT_DIR / filename

    doc.save(str(filepath))
    return str(filepath)
